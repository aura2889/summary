# -*- coding: utf-8 -*-
"""Өндіріс ақпары: месячный Excel (лист на день) -> dict по дням."""
import re, openpyxl
from openpyxl.utils import column_index_from_string as ci

AKROWS=[(6,14),(15,26),(27,37),(38,50),(51,67),(68,75),(77,88),(89,111)]
TOTAL=114
F=dict(own='D',own0='C',buy_p='E',buy_f='F',rent='G',line='K',rep='L',res='M',wo='N',nd='O',
       dp='S',dpd='T',dpr='U',df='V',np='X',npd='Y',npr='Z',nf='AA',
       plan='AW',pd='AX',pr='AY',fact='BB',fd='BC',fr='BD',dev='BG',devd='BE',devr='BF')
CAUSE=dict(rep='BH',repr='BI',staff='BJ',staffr='BK',canc='BL',cancr='BM',yr='BN',yrr='BO',oop='BP')
STAFF=dict(shtat='BQ',list='BR',need='BS',vahta='BT',vac='BU',vacp='BV',vac0='BW',sick='BX',sick5='BY',vacancy='BZ',oth='CA',frm='CB')

def v(ws,r,col):
    x=ws[col+str(r)].value
    if isinstance(x,(int,float)) and not isinstance(x,bool):
        return round(x,4) if isinstance(x,float) else x
    return 0
def nm(s):
    s=str(s or '').split(' / ')[0].strip()
    return re.sub(r'\s+',' ',s)
def row(ws,r):
    d={k:v(ws,r,c) for k,c in F.items()}
    d['c']={k:v(ws,r,c) for k,c in CAUSE.items() if v(ws,r,c)}
    d['s']={k:v(ws,r,c) for k,c in STAFF.items() if v(ws,r,c)}
    kip=ws['CI'+str(r)].value; ktg=ws['CJ'+str(r)].value
    d['kip']=round(kip,3) if isinstance(kip,(int,float)) else None
    d['ktg']=round(ktg,3) if isinstance(ktg,(int,float)) else None
    return d
def day(ws):
    aks=[]
    for i,(h,s) in enumerate(AKROWS):
        t=row(ws,s); t['n']=i+1; t['name']=nm(ws['B'+str(h)].value)
        ty=[]
        for r in range(h+1,s):
            a=ws['A'+str(r)].value; b=ws['B'+str(r)].value
            if b is None: continue
            x=row(ws,r)
            if not any(x[k] for k in ('own','plan','fact','line','rep','res','pd')) and not x['dev']: continue
            x['p']=str(a).strip() if a is not None else ''
            x['name']=nm(b)
            ty.append(x)
        t['types']=ty
        aks.append(t)
    tot=row(ws,TOTAL)
    use=[]
    for r in range(143,148):
        b=ws['B'+str(r)].value
        if b: x=row(ws,r); x['name']=nm(b); use.append(x)
    cust={}
    for lab,col in (('ПУ','V'),('КРС','AE'),('СОС','AL')):
        cust[lab]=ws[col+'165'].value if isinstance(ws[col+'165'].value,(int,float)) else 0
    return dict(aks=aks,tot=tot,use=use,cust=cust)
def parse(path):
    wb=openpyxl.load_workbook(path,data_only=True)
    m=re.search(r'(\d{2})\.(\d{4})',path.split('/')[-1])
    mon=(m.group(2)+'-'+m.group(1)) if m else None
    out={}
    for ws in wb.worksheets:
        if not re.fullmatch(r'\d{1,2}',ws.title): continue
        if not v(ws,TOTAL,'AW'): continue
        out[mon+'-'+ws.title.zfill(2)]=day(ws)
    return out
if __name__=='__main__':
    import sys,json
    o=parse(sys.argv[1]); print(len(o),sorted(o)[:2],sorted(o)[-1])
    d=o[sorted(o)[-1]]; print(json.dumps(d['tot'],ensure_ascii=False)[:400]); print([ (a['n'],a['plan'],a['fact'],a['dev'],len(a['types'])) for a in d['aks']])

def publish(root):
    """prod/<дата>.json -> prod/index.json + prod.json (последний день полностью + динамика по дням)"""
    import os, json
    pd = os.path.join(root, 'prod')
    ds = sorted(f[:-5] for f in os.listdir(pd) if re.fullmatch(r'\d{4}-\d{2}-\d{2}\.json', f))
    if not ds: return
    tr = []
    for d in ds:
        x = json.load(open(os.path.join(pd, d + '.json'), encoding='utf-8'))
        t = x['tot']
        tr.append(dict(d=d, plan=t['plan'], fact=t['fact'], dev=t['dev'], c=t['c'], ak=[[a['plan'], a['fact'], a['dev']] for a in x['aks']]))
    json.dump({'dates': ds, 'latest': ds[-1]}, open(os.path.join(pd, 'index.json'), 'w'), separators=(',', ':'))
    last = json.load(open(os.path.join(pd, ds[-1] + '.json'), encoding='utf-8'))
    last['date'] = ds[-1]; last['trend'] = tr
    json.dump(last, open(os.path.join(root, 'prod.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
