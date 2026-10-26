import pickle, collections, sys, openpyxl
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
DATE_STR=sys.argv[1] if len(sys.argv)>1 else '05.10.2026'
ALL,EXC,NOTES=pickle.load(open('res.pkl','rb'))
FINAL={'ПЧ-Семсер':'ОБиОТ','ПЧ':'ОБиОТ','ЧАКУР мед':'ОБиОТ','ЧАКУР мед.':'ОБиОТ','УПТОК':'БМТС','ГУ':'ГЦ','БПО 2':'ЦПКРС-2','РМУ 2':'Гараж-2','ЦИТС.':'ЦИТС','Газ.цех':'ГЦ','РСУ':'РСЦ','Экология':'ГООС','оппоIR':'IR'}
for _x in ALL: _x['dept']=FINAL.get(_x['dept'],_x['dept'])
for _i in range(6): ALL.append(dict(cell='—',src='MAN',srcname='Добавлено вручную (пожарные машины)',hdr='ПЧ',dept='ОБиОТ',plate='Пожарная машина %d'%(_i+1),code='10',day=1 if _i<3 else 0,night=0 if _i<3 else 1,note='пожарные: 3 день + 3 ночь (по указанию, в диспетчерских файлах нет)',raw='',pid=''))
TPLF="resources/template_first_day.xlsm"
TPL=openpyxl.load_workbook(TPLF).active
tpl_types=[(str(TPL.cell(r,1).value),TPL.cell(r,2).value) for r in range(3,47)]
TN={'10':'Пожарная машина','11':'Скорая','12':'Toyota LC','12.1':'Toyota LC самовождение','13':'Toyota/Jac пикап','13.1':'Toyota/Jac пикап самовождение','16':'Бортовой 1,5т','18':'Камаз, Маз 6-10т','19':'Микроавтобус','20':'Автобус','21':'Сварочные агрегаты','22':'Седельный тягач','22.1':'Эвакуатор','23':'Седельный тягач с лебедкой','24':'Платформа с лебедкой','27':'Манипулятор','28':'Кран 16-20т','29':'Кран 30т','30':'Кран 50т','31':'ЦА-320','32':'Самосвал','33':'Компрессор','34':'ППУА','35':'АЦН','35.1':'Сед.тягач с прицепом АЦН 30т','35.2':'Сед.тягач без прицепа АЦН 30т','36':'Питьевая АЦН','37':'Бензовоз, масловоз','38':'АГП','40':'Мусоровоз','41':'Термобудка','42':'Глубинная лебедка','42.1':'АИС для иссл.скв. с приц','43':'Траншеекопатель','44':'Анкерный Ямобур','45':'Ямобур','46':'Экскаватор','47':'Тракторный экскаватор','48':'Лаборатория','49':'Вакуумная установка','50':'Бульдозер','51':'Ассенизатор','52':'Трубоукладчик','53':'Фронт.погрузчик / Komatsu FD30T-17','54':'Автопогрузчик 1,5т','55':'Погрузчик МКСМ-0,5т','56':'Грейдер','57':'Трактор колесный','57.1':'Трактор колесн. с прицепом','59':'Сваб (подъемник)','60':'Каток','62':'АРОК','63':'Автокран 100тн'}

for c,n in tpl_types: TN[c]=n
TN['35.2']='Сед.тягач без прицепа АЦН 30т'
TN['46']='Гусеничный Экскаватор'
TPL_D=['ЦДН-1','ЦДН-2','ЦДН-3','ЦППН','ЦППД','ЭЦ','ОКС','РСЦ','ЦИТС','ПРЦЭО','ГЦ','ЦПП','БМТС','ЦПКРС-1','ЦПКРС-2','Гараж-1','ЦНИПР','ОБиПКРС','ЦАП','ГООС','Ақкетік','Нейрон','Бурение','ЦТВС','ТБ','ОТДН','ПЧ-Семсер','ПРЦЭО(2)','БД','ОГиР','ГСРиВК','МПГУ','БПО 2','ЛД','Отделкадр','ЦИТС.','ОБО','УЖЭО','Мухтар/Хамит','ОБО-    СБ','РМУ 2','ГАПиИТ','ОМ','оппоIR','Маркшейдер','ПСО']

# план: общий (по типам) и по цехам (день/ночь) - из resources/plan_dept.json (из файла диспетчера, зелёные столбцы)
import json as _json
PD=_json.load(open('resources/plan_dept.json',encoding='utf-8'))
PLAN=collections.Counter({k:v for k,v in PD['total'].items()})
codes_data=collections.Counter()
for x in ALL: codes_data[x['code']]+=x['day']+x['night']
types=[c for c,_ in tpl_types]
types+=[c for c,_ in codes_data.most_common() if c not in types]
types+=[c for c in ['42.1','22.1','63','19','29','41','43'] if c not in types]
for _c in [c for c in PLAN if c not in types]:
    _b=_c.split('.')[0]
    if _b in types: types.insert(types.index(_b)+1,_c)
    else: types.append(_c)
types.remove('23'); types.insert(types.index('50')+1,'23')
types[14:]=sorted(types[14:],key=lambda c:float(c))   # строки 17 и ниже - по возрастанию номера позиции
dvol=collections.Counter()
for x in ALL: dvol[x['dept']]+=x['day']+x['night']
depts=[d for d in PD['depts'] if d!='ОБиОТ']+[d for d,_ in dvol.most_common() if d not in PD['depts'] and d!='ОБиОТ']+['ОБиОТ']
PRI={'34','47','49','28','22','27','31','35','50'}
thin=Side(style='thin'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
YEL=PatternFill('solid',fgColor='FFFF00'); ORG=PatternFill('solid',fgColor='FFC000'); GRY=PatternFill('solid',fgColor='F2F2F2')
NF='0;\\-0;;@'
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Выход ТС'
dws=wb.create_sheet('Ежедневная сводка'); dws=wb.create_sheet('Исходные данные')
hd=['Ячейка','Файл','Шапка в файле','Отдел','Номер','Префикс','Название типа','Лист / дата','День','Ночь','Примечание','Источник УТТ']
for i,h in enumerate(hd,1):
    c=dws.cell(1,i,h); c.font=Font(name='Arial',bold=True,size=10); c.fill=PatternFill('solid',fgColor='D9D9D9'); c.border=B
UTT_SRC={'K8','D7','PKRS'}
rows=sorted(ALL,key=lambda x:(x['src'],x['dept'],int(''.join(ch for ch in x['cell'] if ch.isdigit()) or 0),x['cell']))
for r,x in enumerate(rows,2):
    vals=[x['cell'],x['srcname'],x['hdr'],x['dept'],x['plate'],x['code'],TN.get(x['code'],''),DATE_STR,x['day'],x['night'],x['note'],'УТТ' if x['src'] in UTT_SRC else '']
    for i,v in enumerate(vals,1):
        c=dws.cell(r,i,v); c.font=Font(name='Arial',size=10)
        if i in (6,9,10): c.alignment=Alignment(horizontal='center')
        if i==6: c.number_format='@'
N=len(rows)+1
for col,w in zip('ABCDEFGHIJKL',[8,34,22,14,26,9,34,12,7,7,46,10]): dws.column_dimensions[col].width=w
dws.freeze_panes='A2'; dws.auto_filter.ref=f"A1:L{N}"
RNG=lambda col:f"'Исходные данные'!${col}$2:${col}${N+300}"
def sty(c,fill=None,bold=False,center=False):
    c.font=Font(name='Arial',size=9,bold=bold); c.border=B; c.number_format=NF
    if fill: c.fill=fill
    c.alignment=Alignment(horizontal='center',vertical='center')
    if center: c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
D0=9          # цеха начинаются с I: на каждый цех 4 столбца = план Д, план Н (зелёные, сгруппированы и скрыты), факт Д, факт Н
GRN=PatternFill('solid',fgColor='FF92D050')
for col,t in zip((1,2,3,4,5,6,7,8),('Префикс','Название типа','План','По формуле','Факт','Откл','Всего Д','Всего Н')):
    ws.cell(1,col,t); ws.merge_cells(start_row=1,start_column=col,end_row=2,end_column=col)
for r in (1,2):
    for c in range(1,D0): sty(ws.cell(r,c),YEL if c<3 else ORG,True,True)
for i,d in enumerate(depts):
    c1=D0+4*i
    ws.cell(1,c1,'План'); ws.merge_cells(start_row=1,start_column=c1,end_row=1,end_column=c1+1)
    ws.cell(1,c1+2,d); ws.merge_cells(start_row=1,start_column=c1+2,end_row=1,end_column=c1+3)
    for k,t in enumerate(('Д','Н','Д','Н')): ws.cell(2,c1+k,t)
    for c in (c1,c1+1): sty(ws.cell(1,c),GRN,True,True); sty(ws.cell(2,c),GRN,True,True)
    for c in (c1+2,c1+3): sty(ws.cell(1,c),YEL,True,True); sty(ws.cell(2,c),YEL,True,True)
r0=3; last_col=D0-1+4*len(depts); LC=L(last_col)
PL=[L(D0+4*i+k) for i in range(len(depts)) for k in (0,1)]      # столбцы плана
FDc=[L(D0+4*i+2) for i in range(len(depts))]; FNc=[L(D0+4*i+3) for i in range(len(depts))]   # факт день / ночь
for k,code in enumerate(types):
    r=r0+k; pri=code in PRI; b28=(code=='28')
    ws.cell(r,1,code).number_format='@'; ws.cell(r,2,TN[code])
    sty(ws.cell(r,1),YEL if pri else None,True); sty(ws.cell(r,2),YEL if pri else None,b28)
    ws.cell(r,3,PLAN.get(code,0))
    ws.cell(r,4,'='+'+'.join(f'{c}{r}' for c in PL))
    ws.cell(r,5,f"=G{r}+H{r}"); ws.cell(r,6,f"=E{r}-C{r}")
    ws.cell(r,7,'='+'+'.join(f'{c}{r}' for c in FDc)); ws.cell(r,8,'='+'+'.join(f'{c}{r}' for c in FNc))
    for c in (3,4,5,6,7,8): sty(ws.cell(r,c),ORG,b28)
    ws.cell(r,1).number_format='@'
    for i,d in enumerate(depts):
        c1=D0+4*i; dl=L(c1+2)
        pd=(PD['plan'].get(code) or {}).get(d) or [0,0]
        ws.cell(r,c1,pd[0]); ws.cell(r,c1+1,pd[1])
        ws.cell(r,c1+2,f"=SUMIFS({RNG('I')},{RNG('D')},{dl}$1,{RNG('F')},$A{r})")
        ws.cell(r,c1+3,f"=SUMIFS({RNG('J')},{RNG('D')},{dl}$1,{RNG('F')},$A{r})")
        for c in (c1,c1+1): sty(ws.cell(r,c),GRN,False)
        for c in (c1+2,c1+3): sty(ws.cell(r,c),YEL if pri else None,b28)
tr=r0+len(types)
ws.cell(tr,2,'ИТОГО'); sty(ws.cell(tr,1),GRY,True); sty(ws.cell(tr,2),GRY,True)
for c in range(3,last_col+1):
    cl=L(c); ws.cell(tr,c,f"=SUM({cl}{r0}:{cl}{tr-1})"); sty(ws.cell(tr,c),ORG if c<D0 else GRY,True)
ws.cell(tr+1,2,'ВСЕГО день+ночь'); sty(ws.cell(tr+1,1),GRY,True); sty(ws.cell(tr+1,2),GRY,True)
for c in range(3,D0): sty(ws.cell(tr+1,c),ORG,True)
ws.cell(tr+1,4,'='+'+'.join(f'{c}{tr}' for c in PL)); ws.cell(tr+1,5,f"=E{tr}")
for i in range(len(depts)):
    c1=D0+4*i
    for q in (c1,c1+2):
        ws.cell(tr+1,q,f"={L(q)}{tr}+{L(q+1)}{tr}"); ws.merge_cells(start_row=tr+1,start_column=q,end_row=tr+1,end_column=q+1)
        for c in (q,q+1): sty(ws.cell(tr+1,c),GRY,True,True)
ws.column_dimensions['A'].width=8; ws.column_dimensions['B'].width=30
for c in range(3,last_col+1): ws.column_dimensions[L(c)].width=7
ws.column_dimensions['D'].hidden=True
for i in range(len(depts)):                     # зелёные столбцы плана каждого цеха - группа «+», по умолчанию свёрнута
    c1=D0+4*i; ws.column_dimensions.group(L(c1),L(c1+1),outline_level=1,hidden=True)
ws.row_dimensions[1].height=30; ws.freeze_panes='C3'
# short sheet from user's formatted template
tsrc=openpyxl.load_workbook('resources/short_template_user.xlsx')['Ежедневная сводка']
sh=wb['Ежедневная сводка']
for row in tsrc.iter_rows():
    for c in row:
        n=sh.cell(c.row,c.column,c.value)
        if c.has_style:
            n.font=copy(c.font); n.fill=copy(c.fill); n.border=copy(c.border); n.alignment=copy(c.alignment); n.number_format=c.number_format; n.protection=copy(c.protection)
for m in tsrc.merged_cells.ranges: sh.merge_cells(str(m))
for k,v in tsrc.column_dimensions.items():
    sh.column_dimensions[k].width=v.width; sh.column_dimensions[k].min=v.min; sh.column_dimensions[k].max=v.max
for k,v in tsrc.row_dimensions.items():
    if v.height: sh.row_dimensions[k].height=v.height
sh.sheet_view.showGridLines=tsrc.sheet_view.showGridLines
LASTD=max(c.column for c in tsrc[3] if c.value is not None and c.column>=11)
FIRSTD=[c.column for c in tsrc[3] if c.value=='ЦДН-1'][0]
colI={c.value:c.column for c in tsrc[3] if c.value in ('План','Факт','Откл','День','Ночь')}
PC=[c.column for c in tsrc[3] if c.value=='Техника'][0]-1
PCL=L(PC)
for c in tsrc[1]:
    if c.value and str(c.value).strip().startswith('күні'): sh.cell(1,c.column).value='күні                                 '+DATE_STR+'ж.'
for r_ in range(5,24,2):
    if sh.cell(r_,PC).value is None: continue
    for c_ in range(FIRSTD,LASTD+1):
        cl=L(c_)
        mk=lambda col:f"SUMIFS({RNG(col)},{RNG('D')},{cl}$3,{RNG('F')},${PCL}{r_}&\"\")"
        d=mk('I'); n=mk('J')
        sh.cell(r_,c_,f'=IF({d}+{n}=0,"",IF({n}=0,{d},IF({d}=0,"0+"&{n}&"="&{n},{d}&"+"&{n}&"="&({d}+{n}))))')
        sh.cell(r_,c_).number_format='General'
sh.page_setup.orientation=tsrc.page_setup.orientation; sh.page_setup.paperSize=tsrc.page_setup.paperSize
sh.sheet_properties.pageSetUpPr=copy(tsrc.sheet_properties.pageSetUpPr)
sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=1
sh.page_margins=copy(tsrc.page_margins)
sh.print_options.horizontalCentered=tsrc.print_options.horizontalCentered; sh.print_options.verticalCentered=tsrc.print_options.verticalCentered
sh.sheet_view.zoomScale=tsrc.sheet_view.zoomScale
GREEN=Font(name='Times New Roman',size=16,color='FF00B050')
for r_ in range(4,23,2):
    if sh.cell(r_+1,PC).value is None: continue
    for c_ in range(FIRSTD,LASTD+1):
        cl=L(c_)
        mk=lambda col:f"SUMIFS({RNG(col)},{RNG('D')},{cl}$3,{RNG('F')},${PCL}{r_+1}&\"\",{RNG('L')},\"УТТ\")"
        t=f"{mk('I')}+{mk('J')}"
        cell=sh.cell(r_,c_,f'=IF({t}=0,"","утт-"&({t}))')
        cell.font=GREEN; cell.alignment=Alignment(horizontal='center',vertical='center'); cell.number_format='General'
for r_ in range(5,24,2):
    if sh.cell(r_,PC).value is None: continue
    for nm,src in (('План','C'),('Факт','E'),('Откл','F'),('День','G'),('Ночь','H')):
        f=f"INDEX('Выход ТС'!${src}:${src},MATCH(${PCL}{r_}&\"\",'Выход ТС'!$A:$A,0))"
        sh.cell(r_,colI[nm],f'=IFERROR({f},"")' if nm in ('План','Факт','Откл') else f'=IFERROR(IF({f}=0,"",{f}),"")')
        sh.cell(r_,colI[nm]).number_format='General'

# --- отклонение от плана по цехам (факт - план), столбцы справа от таблицы: AE, AF, ... ; в 3-й строке название цеха
_codes=[str(sh.cell(r_,PC).value) for r_ in range(5,24,2) if sh.cell(r_,PC).value is not None]
_fact=collections.Counter()
for x in ALL: _fact[(x['dept'],x['code'])]+=x['day']+x['night']
DV=[d for d in depts if any(sum((PD['plan'].get(c) or {}).get(d) or [0,0]) or _fact[(d,c)] for c in _codes)]
DV0=LASTD+1
_hs=sh.cell(3,LASTD); _bs=sh.cell(5,LASTD)
sh.cell(2,DV0,'Откл. по цехам (факт − план)'); sh.cell(2,DV0).font=Font(name='Times New Roman',size=12,bold=True)
sh.cell(2,DV0).alignment=Alignment(horizontal='left',vertical='center')
for j,d in enumerate(DV):
    c_=DV0+j; cl=L(c_)
    h=sh.cell(3,c_,d)
    h.font=Font(name='Times New Roman',size=11,bold=True); h.fill=copy(_hs.fill); h.border=copy(_hs.border); h.alignment=Alignment(horizontal='center',vertical='center',shrink_to_fit=True)
    sh.column_dimensions[cl].width=9
    i=depts.index(d); c1=D0+4*i
    for r_ in range(5,24,2):
        if sh.cell(r_,PC).value is None: continue
        m=f"MATCH(${PCL}{r_}&\"\",'Выход ТС'!$A:$A,0)"
        ix=lambda col:f"INDEX('Выход ТС'!${L(col)}:${L(col)},{m})"
        cell=sh.cell(r_,c_,f'=IFERROR({ix(c1+2)}+{ix(c1+3)}-{ix(c1)}-{ix(c1+1)},"")')
        cell.font=Font(name='Times New Roman',size=14,bold=True); cell.border=copy(_bs.border)
        cell.alignment=Alignment(horizontal='center',vertical='center'); cell.number_format='[Red]-0;[Color10]+0;;'
sh.print_area=f"A1:{L(LASTD)}{sh.max_row}"       # печать как раньше, без новых столбцов

# --- перечень госномеров по технике: слева пары (B,C) и (E,F), справа с 47-й строки пары через столбец
PLIST=[str(sh.cell(r_,PC).value) for r_ in range(5,24,2) if sh.cell(r_,PC).value is not None]
dord={d:i for i,d in enumerate(depts)}
items=sorted([x for x in ALL if x['code'] in PLIST and (x['day'] or x['night'])],key=lambda x:(PLIST.index(x['code']),dord.get(x['dept'],99),x['plate']))
# очистить зону списка (значения и заливки), выровнять по левому краю
LF=PatternFill(fill_type=None)
zones=[(3,68,2,7),(47,68,9,30)]
for r1_,r2_,c1_,c2_ in zones:
    for r_ in range(r1_,r2_+1):
        for c_ in range(c1_,c2_+1):
            cell=sh.cell(r_,c_); cell.value=None; cell.fill=LF
            a=copy(cell.alignment); a.horizontal='left'; a.vertical='center'; a.shrink_to_fit=False; cell.alignment=a
cols=[]   # (первая строка, последняя строка, столбец номера)
for c0 in (2,5): cols.append((3,68,c0))
for c0 in (15,18,21,24,27): cols.append((47,68,c0))
FILL=PatternFill('solid',fgColor='EBF1DE')
HF=Font(name='Times New Roman',size=14,bold=True)
YEL=PatternFill('solid',fgColor='FFFF00')
LBL={str(sh.cell(r2,PC).value):str(sh.cell(r2,PC+1).value).strip() for r2 in range(5,24,2) if sh.cell(r2,PC).value is not None and sh.cell(r2,PC+1).value}
print(LBL)
pos=0; placed=0; ci=0; r_=cols[0][0]; prevg=None; gi=-1
def nextcol():
    global ci,r_
    ci+=1
    if ci<len(cols): r_=cols[ci][0]
for it in items:
    g=it['code']
    new_group = g!=prevg
    if new_group:
        gi+=1
        if prevg is not None: r_+=1          # пустая строка между префиксами
        # заголовок группы + хотя бы один номер должны поместиться в одном столбце
        if ci<len(cols) and r_+1>cols[ci][1]: nextcol()
        if ci>=len(cols): break
        c0=cols[ci][2]
        for k in (0,1):
            cell=sh.cell(r_,c0+k)
            cell.font=HF
            cell.fill=YEL
        sh.cell(r_,c0,LBL.get(g,TN.get(g,g)))
        r_+=1
    prevg=g
    if ci<len(cols) and r_>cols[ci][1]:
        nextcol()
        if ci<len(cols) and not new_group:
            for k in (0,1):
                cc=sh.cell(r_,cols[ci][2]+k); cc.font=HF
                cc.fill=YEL
            sh.cell(r_,cols[ci][2],LBL.get(g,TN.get(g,g))+' (прод.)'); r_+=1
    if ci>=len(cols): break
    c0=cols[ci][2]
    for k,v in ((0,it['plate']),(1,it['dept'])):
        cell=sh.cell(r_,c0+k,v)
        if gi%2==0: cell.fill=FILL
    r_+=1; placed+=1
# ширина: столбцы номеров шире, разделители (D, G) узкие
from openpyxl.worksheet.dimensions import ColumnDimension
for _k in list(sh.column_dimensions.keys()):
    if _k=='B': del sh.column_dimensions[_k]
for _col,_w in (('B',18),('C',12),('D',4),('E',18),('F',12),('G',4)):
    sh.column_dimensions[_col]=ColumnDimension(sh,index=_col,width=_w,customWidth=True)
print('LIST',len(items),'placed',placed)
LISTINFO=(len(items),placed)
wb.save(f'out/Свод_Диспетчер_{DATE_STR}.xlsx')
print(len(types),len(depts),tr)
