import sys, re, json, openpyxl
from openpyxl.utils import get_column_letter as L

def num(x):
    if x is None or x == '': return 0
    return int(x) if float(x) == int(float(x)) else float(x)

# Названия вкладок «диспетчеры»: источник (столбец B листа «Исходные данные») -> имя
DISP_MAP = [('Автобусы', 'Диспетчер 4'), ('Ақпар', 'Диспетчер 1'), ('Диспетчер №3', 'Диспетчер 3'),
            ('Сводка от', 'Диспетчер 2'), ('Диспетчер №7', 'Диспетчер 7'), ('К-8', 'Диспетчер 8'),
            ('Мобил', 'Мобил Сервис')]
def disp_name(src):
    if 'вручную' in src.lower(): return None
    for key, name in DISP_MAP:
        if src.startswith(key): return name
    return src                                      # новый, неизвестный файл - показываем как есть
def disp_key(name):
    m = re.search(r'(\d+)$', name)
    return (0, int(m.group(1))) if m and name.startswith('Диспетчер') else (1, 0)

def convert(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    # ---- дата
    d1 = str(wb['Ежедневная сводка']['H1'].value or '')
    m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', d1)
    date = f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
    # ---- машины
    veh, disp_names, disp_src = [], [], []
    for r in wb['Исходные данные'].iter_rows(min_row=2, values_only=True):
        if not r[4]: continue
        nm = re.sub(r'^[СC]\s*(?=\d)', '', str(r[4]).strip())
        v = {'dep': str(r[3]), 'num': nm, 'p': str(r[5]), 't': str(r[6]), 'd': num(r[8]), 'n': num(r[9])}
        src = str(r[1] or '')
        name = disp_name(src)
        if name:                                    # «Добавлено вручную» (пожарные) в диспетчеры не входит
            if name not in disp_names: disp_names.append(name); disp_src.append(src)
            v['g'] = disp_names.index(name)
        veh.append(v)
    order = sorted(range(len(disp_names)), key=lambda i: disp_key(disp_names[i]))
    remap = {old: new for new, old in enumerate(order)}
    for v in veh:
        if 'g' in v: v['g'] = remap[v['g']]
    disp = [{'name': disp_names[i], 'src': disp_src[i]} for i in order]
    # ---- типы и итоги
    ws = wb['Выход ТС']
    depOrder = [str(ws.cell(1, c).value) for c in range(11, ws.max_column + 1, 4) if ws.cell(1, c).value]       # K, O, S… (перед каждым цехом 2 столбца плана)
    types, totals = [], None
    for r in range(3, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if b == 'ИТОГО':
            totals = {'plan': num(ws.cell(r, 3).value), 'fact': num(ws.cell(r, 5).value), 'diff': num(ws.cell(r, 6).value),
                      'day': num(ws.cell(r, 7).value), 'night': num(ws.cell(r, 8).value)}
            break
        if ws.cell(r, 1).value is None and b is None: continue
        t = {'p': str(ws.cell(r, 1).value), 'name': str(b), 'plan': num(ws.cell(r, 3).value), 'fact': num(ws.cell(r, 5).value),
             'diff': num(ws.cell(r, 6).value), 'day': num(ws.cell(r, 7).value), 'night': num(ws.cell(r, 8).value)}
        dp = {}                                                  # план по цехам: {цех: [день, ночь]} (только ненулевые)
        for i, dn in enumerate(depOrder):
            a, b2 = num(ws.cell(r, 9 + 4 * i).value), num(ws.cell(r, 10 + 4 * i).value)
            if a or b2: dp[dn] = [a, b2]
        if dp: t['dp'] = dp
        types.append(t)
    # ---- режимный лист
    wd = wb['Ежедневная сводка']
    c0 = 15                                   # O
    deps, cols = [], []
    c_end = next((c for c in range(c0, wd.max_column + 1) if str(wd.cell(2, c).value or '').startswith('Откл. по цехам')), wd.max_column + 1)   # с этого столбца - отклонения по цехам, это не режимная таблица
    for c in range(c0, c_end):
        v = wd.cell(3, c).value
        if v: deps.append(str(v)); cols.append(c)
    rows = []
    for r in range(4, wd.max_row + 1):
        name, plan = wd.cell(r, 9).value, wd.cell(r, 10).value
        if isinstance(name, str) and isinstance(plan, (int, float)):
            row = {'p': str(wd.cell(r, 8).value) if wd.cell(r, 8).value is not None else '', 'name': name, 'plan': num(plan), 'fact': num(wd.cell(r, 11).value), 'diff': num(wd.cell(r, 12).value),
                   'day': num(wd.cell(r, 13).value), 'night': num(wd.cell(r, 14).value), 'cells': {}}
            if row['p'] == '23': row['name'] = 'ТРАЛ с лебедкой'
            for dname, c in zip(deps, cols):
                v = wd.cell(r, c).value
                if v in (None, ''): continue
                cell = {}
                if isinstance(v, str):
                    mm = re.fullmatch(r'\s*(\d+\s*\+\s*\d+)\s*=\s*(\d+)\s*', v)
                    if mm: cell = {'v': int(mm.group(2)), 'e': mm.group(1).replace(' ', '')}
                    else: cell = {'v': num(re.sub(r'\D', '', v) or 0)}
                else: cell = {'v': num(v)}
                u = wd.cell(r - 1, c).value
                if isinstance(u, str):
                    mu = re.search(r'утт\D*(\d+)', u)
                    if mu: cell['u'] = int(mu.group(1))
                row['cells'][dname] = cell
            rows.append(row)
    out = {'date': date, 'totals': totals, 'depOrder': depOrder, 'types': types, 'veh': veh, 'disp': disp, 'daily': {'deps': deps, 'rows': rows}}
    add_tral_winch(out)
    return out

if __name__ == '__main__':
    out = convert(sys.argv[1])
    if len(sys.argv) > 2: json.dump(out, open(sys.argv[2], 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    else: print(json.dumps(out, ensure_ascii=False)[:500])


def add_tral_winch(data):
    """в режимный лист добавляем строку «ТРАЛ с лебедкой» (префикс 23) сразу после ТРАЛ (22): план/факт - из основной таблицы, по отделам - из списка машин"""
    D = data.get('daily') or {}
    rows = D.get('rows')
    if rows is None or any(r.get('p') == '23' for r in rows): return
    ty = [t for t in data.get('types', []) if str(t.get('p')) == '23']
    if not ty: return
    t = ty[0]
    cells = {}
    for v in data.get('veh', []):
        if str(v.get('p')) == '23' and v.get('dep') in D.get('deps', []):
            c = cells.setdefault(v['dep'], {'v': 0}); c['v'] += 1
    row = {'p': '23', 'name': 'ТРАЛ с лебедкой', 'plan': t['plan'], 'fact': t['fact'], 'diff': t['diff'], 'day': t['day'], 'night': t['night'], 'cells': cells}
    ix = next((i for i, r in enumerate(rows) if r.get('p') == '22'), None)
    rows.insert(ix + 1 if ix is not None else len(rows), row)
