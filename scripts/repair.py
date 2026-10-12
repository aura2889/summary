"""Ремонт техники по автоколоннам: читает файлы начальников колонн (лист нужной даты)."""
import re, datetime, subprocess, tempfile, os, zipfile, shutil
import openpyxl

PLATE = re.compile(r'^\s*\d{2}(?:\.\d)?\s*[-–]\s*\S')
def clean(s): return re.sub(r'\s+', ' ', str(s)).strip() if s is not None else ''

def to_xlsx(path):
    """Настоящий .xls (старый формат) -> .xlsx через LibreOffice. Возвращает путь, который читает openpyxl."""
    if zipfile.is_zipfile(path) and path.lower().endswith(('.xlsx', '.xlsm')): return path, None
    tmp = tempfile.mkdtemp()
    if zipfile.is_zipfile(path):                 # файл .xls, но внутри уже xlsx
        out = os.path.join(tmp, 'f.xlsx'); shutil.copyfile(path, out); return out, tmp
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', tmp, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    out = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + '.xlsx')
    return out, tmp

def sheet_date(name, ref):
    m = re.search(r'(\d{1,2})\s*[.,]\s*(\d{1,2})(?:\s*[.,]\s*(\d{2,4}))?', name)
    if not m: return None
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = (2000 + int(y) if len(y) == 2 else int(y)) if y else None
    try:
        return datetime.date(y or ref.year, mo, d), y is not None
    except ValueError: return None

def choose_sheet(names, ref):
    """лист на дату ref; если такого нет - самый свежий лист не позже ref"""
    exact, best = None, None
    for n in names:
        r = sheet_date(n, ref)
        if not r: continue
        dt, has_y = r
        if dt == ref: exact = n
        elif dt < ref and (best is None or dt >= best[0]): best = (dt, n)
    if exact: return exact, ref
    if best: return best[1], best[0]
    return None, None

def header_map(row):
    h = {}
    for j, v in enumerate(row):
        t = clean(v).lower()
        if not t: continue
        if 'мемлекеттік' in t or 'госномер' in t: h.setdefault('plate', j)
        elif 'марка' in t: h.setdefault('make', j)
        elif t.startswith('түрі'): h.setdefault('type', j)
        elif 'шыққан жылы' in t or 'год выпуска' in t: h.setdefault('year', j)
        elif 'жөндеуге бастап' in t: h.setdefault('start', j)
        elif 'жалпы жөндеуде' in t: h.setdefault('days', j)
        elif 'жөндеу жұмыстарының' in t or 'ремонтные работы' in t: h.setdefault('work', j)
        elif 'ескертпе' in t or 'примечание' in t: h.setdefault('note', j)
        elif 'шығатын күні' in t: h.setdefault('exit', j)
        elif 'дайын' in t: h.setdefault('ready', j)
        elif 'орнына' in t: h.setdefault('repl', j)
        elif 'шыққан көліктер' in t: h.setdefault('out', j)
    return h

def num(x):
    try: return int(float(str(x).replace(',', '.')))
    except Exception: return None

def isodate(x):
    if isinstance(x, datetime.datetime): return x.date().isoformat()
    if isinstance(x, datetime.date): return x.isoformat()
    return None

def parse_sheet(ws, ref):
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:10]) if any('мемлекеттік' in clean(c).lower() or 'госномер' in clean(c).lower() for c in r)), None)
    if hi is None: return None
    H = header_map(rows[hi]); P = H['plate']
    work = H.get('work', 14)
    body = rows[hi + 2:]
    # колонки со списками номеров (замена / вышли / готовы)
    extra = {H[k] for k in ('exit', 'note') if k in H}
    cnt = {}
    for r in body:
        for j, c in enumerate(r):
            if j > work and j not in extra and isinstance(c, str) and PLATE.match(c): cnt[j] = cnt.get(j, 0) + 1
    pc = sorted(cnt)
    labels = {}
    if len(pc) >= 3: labels = {pc[0]: 'repl', pc[1]: 'out', pc[2]: 'ready'}
    else:
        if len(pc) == 2: labels = {pc[0]: 'out', pc[1]: 'ready'}
        else:
            for j in pc:
                hv = [k for k in ('ready', 'repl', 'out') if H.get(k) == j]
                labels[j] = hv[0] if hv else 'out'
    ak, chief, sec, secTxt = None, '', 'ak', ''
    vehs, lists = [], {'repl': [], 'out': [], 'ready': []}
    def add(k, v):
        v = clean(v)
        if v and v not in lists[k]: lists[k].append(v)
    for r in body:
        pv = r[P] if P < len(r) else None
        isveh = isinstance(pv, str) and PLATE.match(pv) and not clean(pv).lower().startswith('жалпы')
        if isveh:
            def g(k):
                j = H.get(k)
                return r[j] if j is not None and j < len(r) else None
            v = {'num': clean(pv), 'make': clean(g('make')), 'type': clean(g('type')),
                 'year': (str(num(g('year'))) if num(g('year')) and num(g('year')) > 1900 else clean(g('year'))),
                 'from': isodate(g('start')), 'k': num(g('days')), 'work': clean(g('work')),
                 'sec': sec, 'secTxt': secTxt}
            nt = clean(g('note'))
            if nt and nt != '1' and not re.fullmatch(r'\d+', nt): v['note'] = nt
            ex = g('exit'); ex = isodate(ex) or clean(ex)
            if ex: v['exit'] = ex
            for j, k in labels.items():
                if k == 'repl' and j < len(r) and isinstance(r[j], str) and PLATE.match(r[j]): v['repl'] = clean(r[j])
            vehs.append(v)
        else:
            text = ' '.join(clean(c) for c in r if isinstance(c, str) and clean(c))
            m = re.search(r'А\s*/\s*т\s*№\s*(\d+)\s*[-–—]?\s*(.*)', text)
            if m: ak, chief, sec, secTxt = int(m.group(1)), clean(m.group(2)), 'ak', ''
            elif re.search(r'КПБ|Көліктерді пайдалану', text): sec, secTxt = 'oets', ''
            elif re.search(r'ТҚКжЖБ|ТҚжЖБ|ОТОиРТС', text): sec, secTxt = 'oto', ''
            elif 'Есептен шығу' in text: sec, secTxt = 'wo', text
        for j, k in labels.items():
            if j < len(r) and isinstance(r[j], str) and PLATE.match(r[j]):
                if k == 'repl' and isveh: continue
                add(k, r[j])
    for v in vehs:
        v['days'] = (ref - datetime.date.fromisoformat(v['from'])).days if v['from'] else v['k']
    return {'n': ak, 'chief': chief, 'vehs': vehs, **{k: lists[k] for k in lists}}

def read_ak(path, ref):
    """-> (dict АК, список предупреждений)"""
    warn = []
    xp, tmp = to_xlsx(path)
    try:
        wb = openpyxl.load_workbook(xp, read_only=True, data_only=True)
        name, sd = choose_sheet(wb.sheetnames, ref)
        if not name: return None, [f'{os.path.basename(path)}: нет листа с датой']
        res = parse_sheet(wb[name], sd)
        if not res: return None, [f'{os.path.basename(path)}: лист «{name}» без таблицы']
        res['date'] = sd.isoformat(); res['sheet'] = name
        if sd != ref: warn.append(f'АК №{res["n"]}: листа на {ref:%d.%m} нет, взят {sd:%d.%m}')
        bad = [v['num'] for v in res['vehs'] if v['from'] and v['k'] is not None and abs(v['k'] - v['days']) > 1]
        if bad: warn.append(f'АК №{res["n"]}: дни в файле не совпадают с датой начала у {len(bad)} машин (считаю по дате)')
        return res, warn
    finally:
        if tmp: shutil.rmtree(tmp, ignore_errors=True)

def is_repair_file(path):
    """файл похож на отчёт автоколонны: на одном из последних листов есть заголовок «Мемлекеттік белгі»"""
    try:
        xp, tmp = to_xlsx(path)
        try:
            wb = openpyxl.load_workbook(xp, read_only=True, data_only=True)
            for n in reversed(wb.sheetnames[-40:]):
                for i, r in enumerate(wb[n].iter_rows(max_row=5, values_only=True)):
                    if any('мемлекеттік' in clean(c).lower() or 'госномер' in clean(c).lower() for c in r): return True
            return False
        finally:
            if tmp: shutil.rmtree(tmp, ignore_errors=True)
    except Exception:
        return False
