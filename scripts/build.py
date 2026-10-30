"""Робот: берёт файлы из input/, собирает data.json и drivers.json, кладёт в архив, чистит старое.
Запуск из корня репозитория: python3 scripts/build.py"""
import sys, os, re, json, glob, shutil, datetime
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import openpyxl
from svod2json import convert
import build_drivers, make_mobil, repair, svod_flow, make_phonebook

KEEP_DAYS = 365
INP = os.path.join(ROOT, 'input'); ARC = os.path.join(ROOT, 'archive')
summary = []
def say(s): print(s); summary.append(s)

def mobil_for(base, day):
    """основной список Mobil + списки вахт data/mobil_vahta_*.json, действующие в этот день (from..to включительно)"""
    import glob
    out = dict(base)
    for fp in sorted(glob.glob(os.path.join(ROOT, 'data', 'mobil_vahta_*.json'))):
        try: v = json.load(open(fp, encoding='utf-8'))
        except Exception: continue
        if v.get('from', '0000') <= day <= v.get('to', '9999'): out.update(v.get('drivers', {}))
    return out

def classify(path):
    """свод / экспорт / вахта / ремонт автоколонны / None - по содержимому, а не по имени файла"""
    try:
        if make_phonebook.is_phonebook(path): return 'book'
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        names = wb.sheetnames
        if 'Исходные данные' in names and 'Выход ТС' in names: return 'svod'
        if sum(1 for n in names if re.fullmatch(r'\d{1,2}', n)) >= 1 and 'Префикс' in str(wb[[n for n in names if re.fullmatch(r'\d{1,2}', n)][0]]['A3'].value or ''): return 'prod'
        ws = wb[names[0]]
        first = next(ws.iter_rows(max_row=1, values_only=True), ())
        if build_drivers.export_cols(first) and any(str(x or '').startswith('Путевой лист') for x in first): return 'export'
        a1 = next(ws.iter_rows(max_row=1, values_only=True), (None,))[0]
        if 'Мобил Сервис' in names or 'Вахта' in str(a1 or ''): return 'roster'
        if repair.is_repair_file(path): return 'repair'
        role = svod_flow.detect_role(path)
        if role: return 'disp:' + role
    except Exception as e:
        say(f'⚠ не удалось прочитать {os.path.basename(path)}: {e}')
    return None

def load(p, default=None):
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else default
def dump(obj, p):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    json.dump(obj, open(p, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))

def aq_now():
    return datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5)))

def aq_today():
    return datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5))).date()

def publish(data, drv, svod_file):
    """кладёт день в архив, обновляет индекс дат и «сегодня» для сайта"""
    date = data['date']; day = os.path.join(ARC, date)
    dump(data, os.path.join(day, 'data.json'))
    if drv is not None: dump(drv, os.path.join(day, 'drivers.json'))
    if svod_file:
        os.makedirs(day, exist_ok=True)
        if os.path.abspath(svod_file) != os.path.abspath(os.path.join(day, 'svod.xlsx')): shutil.copy(svod_file, os.path.join(day, 'svod.xlsx'))
    dates = sorted((d for d in os.listdir(ARC) if re.fullmatch(r'\d{4}-\d{2}-\d{2}', d) and os.path.exists(os.path.join(ARC, d, 'data.json'))), reverse=True) if os.path.isdir(ARC) else []
    limit = (datetime.date.today() - datetime.timedelta(days=KEEP_DAYS)).isoformat()
    for d in [d for d in dates if d < limit]:
        shutil.rmtree(os.path.join(ARC, d)); say(f'Удалён старый день {d}')
    dates = [d for d in dates if d >= limit]
    dump({'dates': dates, 'latest': dates[0]}, os.path.join(ARC, 'index.json'))
    if dates[0] == date:
        dump(data, os.path.join(ROOT, 'data.json'))
        if drv is not None: dump(drv, os.path.join(ROOT, 'drivers.json'))
        say(f'Сайт обновлён на {date}')
    else:
        say(f'Загружен более старый день {date}; на сайте остаётся {dates[0]}')
    say(f'В архиве дней: {len(dates)}')

def main():
    files = [f for f in glob.glob(os.path.join(INP, '*')) if f.lower().endswith(('.xlsx', '.xlsm', '.xls')) and not os.path.basename(f).startswith('~$')]
    kinds, conv, tmps, junk = {}, {}, [], []
    ign_p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ignore.txt')      # части названий файлов, которые не нужны (по строке)
    ign = [l.strip().lower() for l in open(ign_p, encoding='utf-8') if l.strip() and not l.startswith('#')] if os.path.exists(ign_p) else []
    for f in list(files):
        if any(x in os.path.basename(f).lower() for x in ign):
            say(f'Файл {os.path.basename(f)} в списке ненужных (scripts/ignore.txt) - пропущен'); junk.append(f); files.remove(f)
    for f in files:
        try:
            xp, tmp = repair.to_xlsx(f)          # старые .xls -> .xlsx
        except Exception as e:
            say(f'⚠ не удалось открыть {os.path.basename(f)}: {e}'); continue
        conv[f] = xp
        if tmp: tmps.append(tmp)
        k = classify(xp)
        if k: kinds.setdefault(k, []).append(f)
        else: say(f'⚠ файл {os.path.basename(f)} не распознан - пропущен'); junk.append(f)
    errors = []
    ref = datetime.date.fromisoformat(os.environ['SVOD_DATE']) if os.environ.get('SVOD_DATE') else aq_today()
    today = ref                                            # «сегодня» = дата, на которую собираем (для проверки в тестах можно задать SVOD_DATE)
    state_p = os.path.join(ROOT, 'data', 'state.json'); state = load(state_p, {})
    RAW = os.path.join(ROOT, 'raw'); os.makedirs(RAW, exist_ok=True)
    # --- Мобил: новый список вахты
    mob_path = os.path.join(ROOT, 'data', 'mobil.json')
    mobil = load(mob_path, {})
    for f in kinds.get('roster', []):
        new = make_mobil.read_roster(conv[f]); mobil.update(new); dump(mobil, mob_path)
        say(f'Список вахты Mobil обновлён: {len(new)} машин')
    for f in sorted(kinds.get('book', []), key=os.path.getmtime)[-1:]:      # главный телефонный справочник (КБМ1-7)
        cp = os.path.join(ROOT, 'data', 'contacts.json'); ct = load(cp)
        if ct is None: errors.append('нет data/contacts.json'); continue
        r0, j0 = make_phonebook.build_rows0(conv[f], ROOT)
        ct['rows0'] = r0; ct['junk'] = sorted(set(ct.get('junk', [])) | set(j0)); dump(ct, cp)
        say(f'Справочник телефонов обновлён: {len(r0)} человек')
    # --- файлы диспетчеров -> raw/<РОЛЬ>.xlsx (хранится последний файл роли)
    newroles = []
    for k, fl in sorted(kinds.items()):
        if not k.startswith('disp:'): continue
        role = k[5:]
        if role == 'GAR2':                         # объединённый файл Гаража-2 = диспетчеры 7 и 8 сразу
            fl = sorted(fl, key=os.path.getmtime)
            shutil.copy(conv[fl[-1]], os.path.join(RAW, 'GAR2.xlsx'))
            tt = aq_now().strftime('%Y-%m-%dT%H:%M')
            for r2 in ('D7', 'K8'):
                state[r2] = today.isoformat(); newroles.append(r2); state.setdefault('t', {})[r2] = tt
                state.setdefault('via', {})[r2] = 'GAR2'
            say('Получен объединённый файл Гараж-2: диспетчеры 7 и 8')
            continue
        state.get('via', {}).pop(role, None)       # отдельный файл роли снова пришёл - объединённый для неё не нужен
        def _rank(f):                              # лишние/старые файлы той же роли: берём с листом на сегодня и сохранённый позже всех
            try:
                wb = openpyxl.load_workbook(conv[f], read_only=True)
                fresh_sheet = 1 if svod_flow.pick_sheet(wb.sheetnames, ref)[1] else 0
                mod = wb.properties.modified or datetime.datetime(2000, 1, 1); wb.close()
                return (fresh_sheet, mod.timestamp(), os.path.getmtime(f))
            except Exception: return (0, 0, os.path.getmtime(f))
        fl = sorted(fl, key=_rank, reverse=True)           # лучший кандидат первый
        if len(fl) > 1:                                    # чужой файл похожего вида (напр. «Гараж 2 …») даёт неизвестные отделы - берём первого «чистого»
            for i, f in enumerate(fl):
                try: bad = svod_flow.unknown_items(conv[f], role, ref)
                except Exception: bad = {'?'}
                if not bad: fl.insert(len(fl), fl.pop(i)); break     # чистый - в конец (дальше берётся fl[-1])
            else: fl = fl[::-1]
        else: fl = fl[::-1]
        if len(fl) > 1: say(f'⚠ {svod_flow.ROLE_DISP[role]}: пришло {len(fl)} файла(ов), взят самый новый ({os.path.basename(fl[-1])})')
        shutil.copy(conv[fl[-1]], os.path.join(RAW, role + '.xlsx')); state[role] = today.isoformat(); newroles.append(role)
        state.setdefault('t', {})[role] = aq_now().strftime('%Y-%m-%dT%H:%M')           # когда пришёл файл роли
        say(f'Получен файл: {svod_flow.ROLE_DISP[role]}')
    svods, exports = kinds.get('svod', []), kinds.get('export', [])
    if exports:
        again = state.get('export') == today.isoformat()                         # выгрузка за сегодня уже приходила
        shutil.copy(conv[sorted(exports, key=os.path.getmtime)[-1]], os.path.join(RAW, 'export.xlsx')); state['export'] = today.isoformat()
        if again and state.get('full', {}).get('d') == today.isoformat():
            pr = state.get('re') if state.get('re', {}).get('d') == today.isoformat() else {}
            rl = sorted(set(pr.get('roles', [])) | {'export'})
            state['re'] = {'d': today.isoformat(), 't': aq_now().strftime('%H:%M'), 'n': len(rl), 'roles': rl}
        state.setdefault('t', {})['export'] = aq_now().strftime('%Y-%m-%dT%H:%M')
        say('Получен экспорт путевых листов')
    data, svod_file = None, None
    if svods:                                              # готовый свод (старый способ)
        if len(svods) != 1: errors.append(f'готовых сводов в input/ найдено {len(svods)}, нужен один')
        else:
            data = convert(conv[svods[0]]); svod_file = conv[svods[0]]
            say(f'Свод на {data["date"]}: план {data["totals"]["plan"]}, факт {data["totals"]["fact"]}, машин {len(data["veh"])}')
    elif newroles or os.environ.get('SVOD_REBUILD') or kinds.get('book'):       # свод собирается из файлов диспетчеров (SVOD_REBUILD=1 - пересобрать из уже лежащих raw/ без новых файлов)
        fresh = {r: state.get(r) == ref.isoformat() for r in svod_flow.ROLES}
        tmpd = os.path.join(ROOT, '.svod_tmp')
        via = tuple(r for r in ('D7', 'K8') if state.get('via', {}).get(r) == 'GAR2' and state.get(r) == ref.isoformat() and os.path.exists(os.path.join(RAW, 'GAR2.xlsx')))
        xl, status, warns, errs = svod_flow.build(RAW, ref, fresh, tmpd, via)
        for w in warns: say('⚠ ' + w)
        if errs:
            for e in errs: say('❌ ' + e)
            errors += errs; say('Свод НЕ опубликован, сайт остаётся на прошлых данных. Файлы сохранены, после исправления правил сборка продолжится при следующем файле.')
        else:
            data = convert(xl); svod_file = xl
            stale = {svod_flow.ROLE_DISP[r]: st['date'] for r, st in status.items() if st['state'] == 'stale'}
            absent = [svod_flow.ROLE_DISP[r] for r, st in status.items() if st['state'] == 'absent']
            if stale: data['stale'] = stale
            if absent: data['absent'] = absent
            data['okd'] = [svod_flow.ROLE_DISP[r] for r, st in status.items() if st['state'] == 'ok']          # у кого файл за сегодня (для плашки)
            if not newroles:                                   # пересборка без новых файлов - время последнего обновления не трогаем
                pv = load(os.path.join(ARC, ref.isoformat(), 'data.json')) or {}
                if pv.get('upd'): data['upd'] = pv['upd']
            if svod_flow.LAST_TRAIN: data['tr'] = len(svod_flow.LAST_TRAIN); data['trl'] = svod_flow.LAST_TRAIN       # на обучении (не в выходе)
            data['xl'] = 1
            data['nf'] = len(svod_flow.ROLES) - len(stale) - len(absent)
            nowt = aq_now().strftime('%H:%M'); rd = ref.isoformat()
            if state.get('full', {}).get('d') != rd: state.pop('full', None); state.pop('re', None)      # новый день - счёт заново
            if 'full' not in state:
                if data['nf'] >= len(svod_flow.ROLES): state['full'] = {'d': rd, 't': nowt}                 # момент, когда набрались все 7 файлов
            elif newroles:                                                                                  # дополнительные обновления считаем накопительно: сколько разных файлов обновилось после полного набора
                pr = state.get('re') if state.get('re', {}).get('d') == rd else {}
                rl = sorted(set(pr.get('roles', [])) | set(newroles))
                state['re'] = {'d': rd, 't': nowt, 'n': len(rl), 'roles': rl}                                 # последнее дополнительное обновление после этого
            say(f'Свод на {data["date"]}: план {data["totals"]["plan"]}, факт {data["totals"]["fact"]}, машин {len(data["veh"])}; ' +
                ('все файлы за сегодня' if not stale and not absent else f'частично: нет свежих файлов у {len(stale) + len(absent)} из 7'))
    elif exports:                                          # пришёл только экспорт: обновим водителей на сегодняшних данных
        for p in (os.path.join(ARC, ref.isoformat(), 'data.json'), os.path.join(ROOT, 'data.json')):
            d0 = load(p)
            if d0 and d0.get('date') == ref.isoformat(): data = d0; break
        if not data: say('⚠ экспорт получен, но данных свода за сегодня ещё нет - водители подставятся при сборке свода')
    if data is not None:
        rd0 = ref.isoformat()
        if state.get('export'): data['exp'] = state['export']                   # за какое число последний экспорт путевых листов (8-й файл для плашки)
        if state.get('full', {}).get('d') == rd0 and data.get('date') == rd0:
            data['full'] = state['full']['t']
            if state.get('re', {}).get('d') == rd0: data['re'] = {'t': state['re']['t'], 'n': state['re']['n']}
        if svods or newroles: data['upd'] = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5))).strftime('%Y-%m-%dT%H:%M')   # когда робот в последний раз обновил свод (время Актобе)
        data['xl'] = 1 if (svod_file or os.path.exists(os.path.join(ARC, data['date'], 'svod.xlsx'))) else data.get('xl')
        if not data['xl']: data.pop('xl', None)
        drv = None
        contacts = load(os.path.join(ROOT, 'data', 'contacts.json'))
        exp = os.path.join(RAW, 'export.xlsx')
        if not contacts: errors.append('нет data/contacts.json'); say('❌ нет data/contacts.json')
        elif os.path.exists(exp):
            drv, rep, miss = build_drivers.build(data, exp, contacts, mobil_for(mobil, data['date']))
            for l in rep: say(l)
            if miss: say(f'Без телефона ({len(miss)}): ' + '; '.join(miss))
            if state.get('export') and state['export'] != data['date']:
                drv['date'] = state['export']; say(f'⚠ экспорт путевых листов за {state["export"]} (за сегодня ещё не приходил)')
        else: say('⚠ экспорта путевых листов ещё не было - окошко водителя будет без данных')
        if drv is None:                                    # водителей не пересобирали - оставить уже лежащих за этот день
            old = load(os.path.join(ARC, data['date'], 'drivers.json')) or load(os.path.join(ROOT, 'drivers.json'))
            if old: drv = old
        publish(data, drv, svod_file)
        shutil.rmtree(os.path.join(ROOT, '.svod_tmp'), ignore_errors=True)
    dump(state, state_p)
    # --- ремонт по автоколоннам
    if kinds.get('repair'):
        if data is not None: ref = datetime.date.fromisoformat(data['date'])
        day = os.path.join(ARC, ref.isoformat()); rp = os.path.join(day, 'repair.json')
        cur = load(rp, {'aks': []})
        aks = {a['n']: a for a in cur['aks'] if a.get('n')}
        warns = []
        for f in kinds['repair']:
            try:
                res, w = repair.read_ak(conv[f], ref)
            except Exception as e:
                say(f'⚠ {os.path.basename(f)}: ошибка чтения отчёта автоколонны: {e}'); continue
            warns += w
            if not res or not res.get('n'):
                say(f'⚠ {os.path.basename(f)}: не определён номер автоколонны'); continue
            aks[res['n']] = res
            say(f'АК №{res["n"]}: машин в ремонте {len(res["vehs"])} (лист {res["sheet"]})')
        for w in warns: say('⚠ ' + w)
        # прошлые дни: в файле АК есть листы за каждое число - заполняем ими архив (только листы ровно на эту дату, не копии)
        try:
            for dd in sorted(os.listdir(ARC), reverse=True):
                try: dt = datetime.date.fromisoformat(dd)
                except ValueError: continue
                if dt >= ref or (ref - dt).days > 31 or not os.path.isdir(os.path.join(ARC, dd)): continue
                drp = os.path.join(ARC, dd, 'repair.json')
                dcur = load(drp, {'aks': []}); daks = {a['n']: a for a in dcur['aks'] if a.get('n')}; got = []
                for f in kinds['repair']:
                    try: r2, _ = repair.read_ak(conv[f], dt)
                    except Exception: continue
                    if r2 and r2.get('n') and r2.get('date') == dt.isoformat(): daks[r2['n']] = r2; got.append(r2['n'])
                if got:
                    dump({'date': dd, 'aks': [daks[k] for k in sorted(daks)]}, drp)
                    say(f'Архив {dd}: отчёты АК обновлены для автоколонн ' + ', '.join(str(i) for i in sorted(got)))
        except Exception as e:
            say(f'⚠ не удалось обновить прошлые дни отчётов АК: {e}')
        out = {'date': ref.isoformat(), 'aks': [aks[k] for k in sorted(aks)]}
        if not aks:
            say('Отчётов автоколонн в файлах не найдено - repair.json не трогаю (файл ошибочно принят за отчёт АК)')
        else:
            dump(out, rp)
            say(f'Автоколонн в отчёте на {ref.isoformat()}: {len(aks)} из 8' + ('' if len(aks) == 8 else ' (не хватает: ' + ', '.join(str(i) for i in range(1, 9) if i not in aks) + ')'))
            idx = load(os.path.join(ARC, 'index.json'), {'dates': []})
            if not idx['dates'] or ref.isoformat() >= idx['dates'][0]:
                dump(out, os.path.join(ROOT, 'repair.json'))
    if kinds.get('prod'):
        import prod
        PD = os.path.join(ROOT, 'prod'); os.makedirs(PD, exist_ok=True)
        for f in sorted(kinds['prod'], key=os.path.getmtime):
            try: days = prod.parse(conv[f])
            except Exception as e: say(f'⚠ Өндіріс ақпары: не прочитан {os.path.basename(f)}: {e}'); continue
            for dd, dv in days.items(): dump(dv, os.path.join(PD, dd + '.json'))
            say(f'Өндіріс ақпары: дней в файле {len(days)} ({min(days)} … {max(days)})' if days else 'Өндіріс ақпары: в файле нет заполненных дней')
        prod.publish(ROOT)
    if not (kinds.get('roster') or kinds.get('book') or svods or exports or newroles or kinds.get('prod')) and not kinds.get('repair'):
        say('В input/ нет подходящих файлов - ничего не делаю')
    for t in tmps: shutil.rmtree(t, ignore_errors=True)
    for ks in kinds.values():          # обработанные файлы убираем из папки
        for f in ks: os.remove(f)
    for f in junk:
        try: os.remove(f)
        except OSError: pass
    if errors:
        open(os.path.join(ROOT, '.build_failed'), 'w').write('\n'.join(errors))
        say('❌ Сборка завершена с ошибками: ' + ' | '.join(errors))
    return 0

if __name__ == '__main__':
    rc = main()
    out = os.environ.get('GITHUB_STEP_SUMMARY')
    if out: open(out, 'a', encoding='utf-8').write('\n'.join('- ' + s for s in summary) + '\n')
    sys.exit(rc)
